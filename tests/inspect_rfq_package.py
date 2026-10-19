"""Inspect the actual downloaded fictional RFQ ZIP and Excel attachment contents."""
import io
import json
import sys
import zipfile
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from rfq_workbench.privacy import check_bundle
from rfq_workbench.rfq import HEADERS

data = Path(sys.argv[1]).read_bytes()
fixture = json.loads((Path(__file__).resolve().parent.parent / "demo" / "requirement-001.json").read_text(encoding="utf-8"))
check_bundle(data, fixture["sensitive_markers"])
with zipfile.ZipFile(io.BytesIO(data)) as archive:
    manifest = json.loads(archive.read("dispatch-manifest.json"))
    inspected = []
    for message in manifest["messages"]:
        book = load_workbook(io.BytesIO(archive.read(message["attachment"])))
        sheet = book["Supplier Reply"]
        assert [cell.value for cell in sheet[1]] == HEADERS
        assert sheet.max_row == len(message["lines"]) + 1
        for row_number, line in enumerate(message["lines"], 2):
            values = [cell.value for cell in sheet[row_number]]
            assert values[:7] == [message["reference"], line["id"], line["line_number"], line["description"], line["specification"] or None, line["quantity"], line["unit"]]
            assert all(value is None for value in values[7:])
        assert all(cell.data_type != "f" for row in sheet for cell in row)
        inspected.append({"category": message["category"], "reference": message["reference"], "positions": len(message["lines"]), "bcc": len(message["bcc"])})
        book.close()
print(json.dumps({"passed": True, "inspected_actual_excel_files": inspected}, indent=2))
