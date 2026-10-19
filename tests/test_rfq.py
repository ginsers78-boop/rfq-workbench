import copy
import io
import json
import tempfile
import unittest
import zipfile
from pathlib import Path

from openpyxl import load_workbook

from rfq_workbench.extraction import ExtractionError
from rfq_workbench.privacy import check_bundle, check_text, suggested_markers
from rfq_workbench.rfq import HEADERS, approve_rfq, create_rfq, download_rfq, reply_workbook, save_privacy
from rfq_workbench.store import Store
from rfq_workbench.suppliers import list_suppliers, save_supplier


class RFQTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.store = Store(Path(self.directory.name) / "test.sqlite3", [{"id": "Fasteners"}, {"id": "Electrical"}, {"id": "Unclassified"}])
        source = "Customer: Fictional Harbor Trading\nCustomer reference: PRIVATE-CUSTOMER-442\nSelling price: confidential 999 OMR\n1. Bolt 12 pcs\n2. Cable 200 m\n3. removed product"
        lines = [{"line_number": "1", "description": "Bolt", "specification": "M10", "quantity": 12, "unit": "pcs", "category": "Fasteners", "source_quote": "1. Bolt 12 pcs", "needs_review": True},
                 {"line_number": "2", "description": "Cable", "specification": "4mm2", "quantity": 200, "unit": "m", "category": "Electrical", "source_quote": "2. Cable 200 m", "needs_review": True},
                 {"line_number": "3", "description": "removed product", "specification": None, "quantity": None, "unit": None, "category": "Unclassified", "source_quote": "3. removed product", "needs_review": True}]
        self.record = self.store.create({"name": "PRIVATE-CUSTOMER-442.txt", "text": source}, {"extraction": {"lines": lines}})
        reviewed = copy.deepcopy(self.record["lines"])
        for line in reviewed[:2]:
            line["approved"] = True
        reviewed[2].update(excluded=True, exclusion_reason="Customer removed this position")
        self.record = self.store.review(self.record["id"], {"revision": 1, "lines": reviewed})
        self.record = save_privacy(self.store, self.record["id"], {"revision": 2, "markers": [], "confirmed": True})
        self.first = save_supplier(self.store, {"name": "Fictional Supply A", "email": "qa-a@example.com", "categories": ["Fasteners", "Electrical"]})
        self.second = save_supplier(self.store, {"name": "Fictional Supply B", "email": "qa-b@example.com", "categories": ["Fasteners"]})
        self.payload = {"revision": 3, "selection": {"Fasteners": [self.first["id"], self.second["id"]], "Electrical": [self.first["id"]]}}

    def tearDown(self):
        self.directory.cleanup()

    def test_real_bundle_allowlist_and_workbook_ids(self):
        rfq = create_rfq(self.store, self.record["id"], self.payload)
        with self.assertRaisesRegex(ExtractionError, "approve"):
            download_rfq(self.store, rfq["id"])
        approved = approve_rfq(self.store, rfq["id"])
        self.assertEqual(approved["status"], "approved_for_local_export")
        bundle = download_rfq(self.store, rfq["id"])
        check_bundle(bundle, self.record["privacy"]["markers"])
        with zipfile.ZipFile(io.BytesIO(bundle)) as archive:
            manifest = json.loads(archive.read("dispatch-manifest.json"))
            self.assertEqual(len(manifest["messages"]), 2)
            for message in manifest["messages"]:
                self.assertNotIn("source_quote", json.dumps(message))
                self.assertNotIn("PRIVATE-CUSTOMER", message["reference"])
                self.assertEqual(len(message["lines"]), 1)
                line = message["lines"][0]
                self.assertNotEqual(line["line_number"], "3")
                book = load_workbook(io.BytesIO(archive.read(message["attachment"])))
                sheet = book["Supplier Reply"]
                row = dict(zip([cell.value for cell in sheet[1]], [cell.value for cell in sheet[2]]))
                self.assertEqual(list(row), HEADERS)
                self.assertEqual(row["RFQ Reference"], message["reference"])
                self.assertEqual(row["Line ID"], line["id"])
                self.assertEqual(row["Original Line"], line["line_number"])
                self.assertIsNone(row["Unit Price"])
                self.assertTrue(all(cell.data_type != "f" for row in sheet for cell in row))
                book.close()
        self.assertEqual(create_rfq(self.store, self.record["id"], self.payload)["id"], rfq["id"])
        self.assertEqual(len({message["reference"] for message in rfq["messages"]}), 2)

    def test_supplier_stability_duplicate_and_stale_update(self):
        edited = save_supplier(self.store, self.first | {"name": "Updated supplier"})
        self.assertEqual(edited["id"], self.first["id"])
        self.assertEqual(edited["revision"], 2)
        with self.assertRaisesRegex(ExtractionError, "changed"):
            save_supplier(self.store, self.first)
        with self.assertRaisesRegex(ExtractionError, "already exists"):
            save_supplier(self.store, {"name": "Duplicate", "email": "QA-A@EXAMPLE.COM", "categories": ["Fasteners"]})
        self.assertEqual(len(list_suppliers(self.store)), 2)

    def test_supplier_header_injection_and_categories(self):
        for change in [{"email": "a@example.com\r\nBCC: victim@example.com"}, {"categories": ["Unclassified"]}, {"categories": ["invented"]}, {"categories": ["Fasteners", "Fasteners"]}]:
            with self.subTest(change=change), self.assertRaises(ExtractionError):
                save_supplier(self.store, {"name": "Test", "email": "new@example.com", "categories": ["Fasteners"]} | change)

    def test_inactive_and_wrong_category_rejected(self):
        save_supplier(self.store, self.first | {"active": False})
        with self.assertRaisesRegex(ExtractionError, "inactive"):
            create_rfq(self.store, self.record["id"], self.payload)
        payload = copy.deepcopy(self.payload)
        payload["selection"]["Electrical"] = [self.second["id"]]
        with self.assertRaisesRegex(ExtractionError, "category"):
            create_rfq(self.store, self.record["id"], payload)

    def test_selections_must_cover_exact_categories(self):
        for selection in [{}, {"Fasteners": []}, {"Fasteners": [self.first["id"]]}, {"Fasteners": [self.first["id"]] * 2, "Electrical": [self.first["id"]]}]:
            with self.subTest(selection=selection), self.assertRaises(ExtractionError):
                create_rfq(self.store, self.record["id"], {"revision": 3, "selection": selection})

    def test_customer_leak_in_product_is_blocked(self):
        reviewed = copy.deepcopy(self.record["lines"])
        reviewed[0]["specification"] = "For FICTIOnal - Harbor Trading"
        changed = self.store.review(self.record["id"], {"revision": 3, "lines": reviewed})
        with self.assertRaisesRegex(ExtractionError, "confidential"):
            create_rfq(self.store, changed["id"], self.payload | {"revision": changed["revision"]})

    def test_outdated_requirement_and_supplier_block_download(self):
        rfq = create_rfq(self.store, self.record["id"], self.payload)
        approve_rfq(self.store, rfq["id"])
        save_supplier(self.store, self.first | {"name": "Changed contact"})
        with self.assertRaisesRegex(ExtractionError, "supplier changed"):
            download_rfq(self.store, rfq["id"])
        another = create_rfq(self.store, self.record["id"], self.payload)
        self.assertNotEqual(another["id"], rfq["id"])
        self.store.review(self.record["id"], {"revision": 3, "lines": self.record["lines"]})
        with self.assertRaisesRegex(ExtractionError, "outdated"):
            approve_rfq(self.store, another["id"])

    def test_privacy_confirmed_and_original_markers_cannot_be_removed(self):
        values = self.record["privacy"]["markers"]
        self.assertIn("Fictional Harbor Trading", values)
        self.assertIn("999 OMR", values)
        with self.assertRaises(ExtractionError):
            save_privacy(self.store, self.record["id"], {"revision": 3, "markers": [], "confirmed": False})
        record = save_privacy(self.store, self.record["id"], {"revision": 3, "markers": ["Extra internal secret"], "confirmed": True})
        self.assertIn("PRIVATE-CUSTOMER-442", record["privacy"]["markers"])
        self.assertIn("Extra internal secret", record["privacy"]["markers"])
        with self.assertRaisesRegex(ExtractionError, "changed"):
            save_privacy(self.store, self.record["id"], {"revision": 3, "markers": [], "confirmed": True})

    def test_unreviewed_or_missing_privacy_blocks_rfq(self):
        with self.store.connect() as connection:
            record = copy.deepcopy(self.record)
            record.pop("privacy")
            connection.execute("UPDATE requirements SET payload = ? WHERE id = ?", (json.dumps(record), record["id"]))
        with self.assertRaisesRegex(ExtractionError, "checklist"):
            create_rfq(self.store, record["id"], self.payload)
        record["status"] = "needs_review"
        with self.store.connect() as connection:
            connection.execute("UPDATE requirements SET payload = ? WHERE id = ?", (json.dumps(record), record["id"]))
        with self.assertRaisesRegex(ExtractionError, "Review all"):
            create_rfq(self.store, record["id"], self.payload)

    def test_history_and_delete_cascade(self):
        rfq = create_rfq(self.store, self.record["id"], self.payload)
        history = self.store.history(self.record["id"])
        self.assertEqual([item["event"] for item in history], ["extracted", "review_saved", "privacy_confirmed"])
        self.assertEqual(history[0]["record"]["status"], "needs_review")
        self.assertEqual(history[1]["record"]["lines"][2]["exclusion_reason"], "Customer removed this position")
        restored = Store(self.store.path, [{"id": "Fasteners"}, {"id": "Electrical"}])
        self.assertEqual(restored.history(self.record["id"]), history)
        self.store.delete(self.record["id"])
        with self.store.connect() as connection:
            self.assertEqual(connection.execute("SELECT count(*) FROM requirement_history").fetchone()[0], 0)
            self.assertEqual(connection.execute("SELECT count(*) FROM rfqs").fetchone()[0], 0)

    def test_independent_office_xml_scan_and_formula_text(self):
        line = {"id": "permanent-id", "line_number": "1", "description": '=HYPERLINK("https://example.com")', "specification": "Secret & Co", "quantity": 1, "unit": "pcs"}
        workbook = reply_workbook("RFQ-TEST", [line])
        book = load_workbook(io.BytesIO(workbook))
        self.assertEqual(book["Supplier Reply"]["D2"].data_type, "s")
        book.close()
        output = io.BytesIO()
        with zipfile.ZipFile(output, "w") as archive:
            archive.writestr("reply.xlsx", workbook)
        with self.assertRaisesRegex(ExtractionError, "confidential"):
            check_bundle(output.getvalue(), ["Secret & Co"])
        for value in ["PRIVATE CUSTOMER 442", "private-customer-442", "ＰＲＩＶＡＴＥ-CUSTOMER-442"]:
            with self.subTest(value=value), self.assertRaises(ExtractionError):
                check_text(value, ["PRIVATE-CUSTOMER-442"])
        for value in ["999.000 OMR", "OMR 999.00", "999,00 OMR"]:
            with self.subTest(value=value), self.assertRaises(ExtractionError):
                check_text(value, ["999 OMR"])

    def test_exclusion_requires_reason_and_cannot_be_approved(self):
        for change in [{"exclusion_reason": " "}, {"approved": True}]:
            lines = copy.deepcopy(self.record["lines"])
            lines[2].update(change)
            with self.subTest(change=change), self.assertRaises(ExtractionError):
                self.store.review(self.record["id"], {"revision": 3, "lines": lines})

    def test_repeated_original_numbers_keep_distinct_excel_line_ids(self):
        source="1. Bolt, 12 pcs\n1. Nut, 8 pcs"
        lines=[{"line_number":"1","description":name,"specification":None,"quantity":qty,"unit":"pcs","category":"Fasteners","source_quote":quote,"needs_review":True} for name,qty,quote in [("Bolt",12,"1. Bolt, 12 pcs"),("Nut",8,"1. Nut, 8 pcs")]]
        locations=[{"start":source.index(line["source_quote"]),"end":source.index(line["source_quote"])+len(line["source_quote"])} for line in lines]
        record=self.store.create({"name":"duplicate-number.txt","text":source},{"extraction":{"lines":lines},"line_locations":locations,"parts":[{"part":"1/2"},{"part":"2/2"}]})
        for line in record["lines"]:line["approved"]=True
        record=self.store.review(record["id"],{"revision":1,"lines":record["lines"],"coverage_verified":True})
        record=save_privacy(self.store,record["id"],{"revision":2,"markers":[],"confirmed":True})
        rfq=create_rfq(self.store,record["id"],{"revision":3,"selection":{"Fasteners":[self.first["id"]]}})
        approve_rfq(self.store,rfq["id"])
        with zipfile.ZipFile(io.BytesIO(download_rfq(self.store,rfq["id"]))) as archive:
            book=load_workbook(io.BytesIO(archive.read(rfq["messages"][0]["attachment"])))
            sheet=book["Supplier Reply"]
            self.assertEqual(sheet["C2"].value,"1")
            self.assertEqual(sheet["C3"].value,"1")
            self.assertNotEqual(sheet["B2"].value,sheet["B3"].value)
            book.close()
