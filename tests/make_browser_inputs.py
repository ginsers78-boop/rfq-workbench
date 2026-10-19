"""Generate fictional, real file fixtures for the browser integration check."""
import io
import json
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas
from PIL import Image, ImageDraw

root = Path(__file__).resolve().parent.parent
out = root / "data" / "browser-inputs"
out.mkdir(parents=True, exist_ok=True)
source = json.loads((root / "demo" / "requirement-001.json").read_text(encoding="utf-8"))["source"]
(out / "requirement.txt").write_text(source, encoding="utf-8")
pdf = canvas.Canvas(str(out / "requirement.pdf"))
pdf.setFont("Helvetica", 11)
for index, line in enumerate(source.splitlines()):
    pdf.drawString(40, 800 - index * 22, line)
pdf.save()
doc = Document()
for line in source.splitlines():
    doc.add_paragraph(line)
doc.save(out / "requirement.docx")
book = Workbook()
book.active.title = "Requirement"
for line in source.splitlines():
    book.active.append([line])
book.save(out / "requirement.xlsx")
image = Image.new("RGB", (1000, 600), "white")
draw = ImageDraw.Draw(image)
draw.text((30, 30), "IMAGE ONLY: 1. Brass valve 12 pcs", fill="black", font_size=32)
scan = canvas.Canvas(str(out / "scan.pdf"))
scan.drawImage(ImageReader(image), 40, 400, width=500, height=300)
scan.save()
(out / "broken.pdf").write_bytes(b"%PDF-1.7\ncorrupted")
print(out)
