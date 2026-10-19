"""Long fictional procurement with repeated source numbering and multiline specs."""
import json
import textwrap
from pathlib import Path
from reportlab.pdfgen import canvas
from openpyxl import Workbook
from docx import Document

root=Path(__file__).resolve().parent.parent
out=root/"data"/"long-inputs"
out.mkdir(parents=True,exist_ok=True)
profiles=[("SS hex bolt M10x50, grade A2-70",120,"pcs","Fasteners","Thread tolerance: 6g"),
          ("Insulated copper wire 2.5mm2, blue",350,"m","Electrical","Conductor class: 5"),
          ("PVC elbow 90 degree, 25mm",40,"pcs","Pipes and fittings","Pressure rating: PN16"),
          ("Stainless steel flat washers M8",500,"pcs","Fasteners","Finish: plain"),
          ("Brass ball valve 1/2 inch BSP",12,"pcs","Pipes and fittings","Seal material: PTFE"),
          ("Galvanised screw M6x30",None,None,"Fasteners","Surface treatment: zinc plated")]
notes=("The supplied item must match the stated dimensions and material. Alternatives require a separate technical description and explicit buyer approval. "
       "Keep this position separate from other requested products. Provide manufacturer identification, traceability and the relevant compliance certificate. "
       "Protect the product against corrosion and mechanical damage during transport. Clearly label packages with the requested item specification. "
       "State any deviation from this requirement in the supplier reply rather than silently substituting an equivalent. "
       "Inspection will compare the delivered product with its requested technical specification; submit supporting technical literature with the offer.")
blocks=["Customer: Fictional Longform Trading\nCustomer reference: PRIVATE-LONGFORM-2026\nSelling price: confidential 4567.890 OMR\n\n"]
labels=[]
for index in range(24):
    description,qty,unit,category,tail=profiles[index%len(profiles)]
    number=index%12+1
    text=f"{number}. {description}, "+(f"{qty} {unit}." if qty is not None else "quantity and unit to be confirmed.")+"\n"+notes+f"\n{tail}.\n\n"
    labels.append({"line_number":str(number),"description":description,"specification":tail,"quantity":qty,"unit":unit,"category":category,"tail":tail.split(": ")[1],"quote":text.splitlines()[0]})
    blocks.append(text)
source="".join(blocks)
(out/"requirement.txt").write_text(source,encoding="utf-8")
(out/"expected.json").write_text(json.dumps({"fictional":True,"expected":labels},indent=2),encoding="utf-8")
pdf=canvas.Canvas(str(out/"requirement.pdf"))
for start in range(0,24,4):
    pdf.setFont("Helvetica",9)
    y=800
    if start==0:
        for line in blocks[0].splitlines():
            pdf.drawString(40,y,line);y-=12
    for block in blocks[start+1:start+5]:
        for paragraph in block.splitlines():
            for line in textwrap.wrap(paragraph,width=110) or [""]:
                pdf.drawString(40,y,line);y-=11
        y-=10
    pdf.showPage()
pdf.save()
book=Workbook()
book.active.title="Requirement"
word=Document()
for block in blocks:
    for line in block.splitlines():
        book.active.append([line])
        word.add_paragraph(line)
book.save(out/"requirement.xlsx")
word.save(out/"requirement.docx")
table=Workbook()
table.remove(table.active)
for sheet_number in range(2):
    sheet=table.create_sheet(f"Category sheet {sheet_number+1}")
    headers=["Line","Description","Specification","Quantity","Unit","Notes"] if sheet_number==0 else ["Line","Unit","Quantity","Specification","Description","Notes"]
    sheet.append(headers)
    for index,(description,qty,unit,category,tail) in enumerate(profiles,1):
        values={"Line":index,"Description":description,"Specification":tail,"Quantity":qty,"Unit":unit,"Notes":notes}
        sheet.append([values[key] for key in headers])
table.save(out/"table-order.xlsx")
print(json.dumps({"directory":str(out),"utf8_bytes":len(source.encode()),"positions":len(labels),"pdf_pages":6}))
