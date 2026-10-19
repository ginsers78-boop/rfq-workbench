import io
import unittest

from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader
from PIL import Image

from rfq_workbench.extraction import ExtractionError
from rfq_workbench.readers import MAX_BYTES, read_document, validate_text


class ReaderTests(unittest.TestCase):
    def test_txt_utf8(self):
        result = read_document("parts.txt", "\ufeff1. Cable 4mm2, 200 m".encode("utf-8"))
        self.assertEqual(result["text"], "1. Cable 4mm2, 200 m")

    def test_pdf_actual_bytes(self):
        buffer = io.BytesIO()
        document = canvas.Canvas(buffer)
        document.drawString(40, 780, "1. Brass ball valve, 12 pcs")
        document.save()
        result = read_document("parts.pdf", buffer.getvalue())
        self.assertIn("Brass ball valve", result["text"])
        self.assertEqual(result["pages"], 1)

    def test_pdf_without_text(self):
        buffer = io.BytesIO()
        writer = PdfWriter()
        writer.add_blank_page(500, 500)
        writer.write(buffer)
        with self.assertRaisesRegex(ExtractionError, "OCR"):
            read_document("scan.pdf", buffer.getvalue())

    def test_encrypted_pdf(self):
        buffer = io.BytesIO()
        writer = PdfWriter()
        writer.add_blank_page(500, 500)
        writer.encrypt("secret")
        writer.write(buffer)
        with self.assertRaisesRegex(ExtractionError, "Password"):
            read_document("private.pdf", buffer.getvalue())

    def test_mixed_pdf_does_not_silently_skip_scanned_page(self):
        buffer = io.BytesIO()
        document = canvas.Canvas(buffer)
        document.drawString(40, 780, "1. Valve 12 pcs")
        document.showPage()
        document.drawImage(ImageReader(Image.new("RGB", (100, 100), "white")), 40, 500)
        document.save()
        with self.assertRaisesRegex(ExtractionError, "page 2.*OCR.*no partial"):
            read_document("mixed.pdf", buffer.getvalue())

    def test_text_header_does_not_hide_scanned_table_on_same_page(self):
        buffer = io.BytesIO()
        document = canvas.Canvas(buffer)
        document.drawString(40, 780, "Enquiry header only")
        document.drawImage(ImageReader(Image.new("RGB", (100, 100), "white")), 40, 500)
        document.save()
        with self.assertRaisesRegex(ExtractionError, "needs OCR"):
            read_document("hybrid.pdf", buffer.getvalue())

    def test_embedded_office_images_are_not_silently_ignored(self):
        from openpyxl.drawing.image import Image as ExcelImage
        image = io.BytesIO()
        Image.new("RGB", (40, 40), "white").save(image, "PNG")
        word = Document()
        word.add_paragraph("Header only")
        word.add_picture(io.BytesIO(image.getvalue()))
        word_output = io.BytesIO()
        word.save(word_output)
        book = Workbook()
        book.active.append(["Header only"])
        book.active.add_image(ExcelImage(io.BytesIO(image.getvalue())))
        book_output = io.BytesIO()
        book.save(book_output)
        for name, data in [("x.docx", word_output.getvalue()), ("x.xlsx", book_output.getvalue())]:
            with self.subTest(name=name), self.assertRaisesRegex(ExtractionError, "Embedded Office images"):
                read_document(name, data)

    def test_page_limit(self):
        buffer = io.BytesIO()
        writer = PdfWriter()
        for _ in range(51):
            writer.add_blank_page(500, 500)
        writer.write(buffer)
        with self.assertRaisesRegex(ExtractionError, "50 pages"):
            read_document("long.pdf", buffer.getvalue())

    def test_docx_preserves_table_order(self):
        buffer = io.BytesIO()
        document = Document()
        document.add_paragraph("Before")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Copper cable"
        table.cell(0, 1).text = "200 m"
        document.add_paragraph("After")
        document.save(buffer)
        self.assertEqual(read_document("parts.docx", buffer.getvalue())["text"], "Before\nCopper cable | 200 m\nAfter")

    def test_xlsx_reads_real_cells(self):
        buffer = io.BytesIO()
        workbook = Workbook()
        workbook.active.append(["Line", "Product", "Qty", "Unit"])
        workbook.active.append([10, "Valve", 12, "pcs"])
        workbook.save(buffer)
        self.assertIn("10 | Valve | 12 | pcs", read_document("parts.xlsx", buffer.getvalue())["text"])

    def test_formulas_not_evaluated(self):
        buffer = io.BytesIO()
        workbook = Workbook()
        workbook.active.append(["=1+2"])
        workbook.save(buffer)
        with self.assertRaisesRegex(ExtractionError, "formulas"):
            read_document("parts.xlsx", buffer.getvalue())

    def test_invalid_formats_and_empty(self):
        for name, data in [("x.exe", b"text"), ("x.pdf", b"text"), ("x.pdf", b"%PDF-broken"), ("x.txt", b" "), ("x.docx", b"PKbroken"), ("x.txt", b"\x00text"), ("x.txt", b"\xff"), ("x.txt", b"")]:
            with self.subTest(name=name, data=data), self.assertRaises(ExtractionError):
                read_document(name, data)

    def test_size_and_utf8_budget(self):
        with self.assertRaisesRegex(ExtractionError, "10 MB"):
            read_document("x.txt", b"a" * (MAX_BYTES + 1))
        with self.assertRaisesRegex(ExtractionError, "1 MiB"):
            validate_text("Ж" * 524289)


if __name__ == "__main__":
    unittest.main()
