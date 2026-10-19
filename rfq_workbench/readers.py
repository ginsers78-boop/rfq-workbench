"""Read actual attachments without executing macros, formulas or embedded content."""
from __future__ import annotations

import io
import zipfile
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

from .extraction import ExtractionError
from .chunks import MAX_DOCUMENT_TEXT

MAX_BYTES = 10 * 1024 * 1024
MAX_TEXT = MAX_DOCUMENT_TEXT
MAX_PAGES = 50


def validate_text(text: str) -> str:
    if not isinstance(text, str) or not text.strip():
        raise ExtractionError("The document has no readable text.")
    if "\x00" in text:
        raise ExtractionError("The document contains unsupported binary text.")
    if len(text.encode("utf-8")) > MAX_TEXT:
        raise ExtractionError("The document exceeds the 1 MiB extracted-text limit. No text was truncated.")
    return text.strip()


def read_document(name: str, content: bytes) -> dict:
    if not content:
        raise ExtractionError("The selected file is empty.")
    if len(content) > MAX_BYTES:
        raise ExtractionError("The maximum file size is 10 MB.")
    extension = Path(name).suffix.casefold()
    pages = None
    details = []
    try:
        if extension == ".txt":
            try:
                text = content.decode("utf-8-sig")
            except UnicodeDecodeError:
                raise ExtractionError("Save the text file using UTF-8 encoding and try again.") from None
        elif extension == ".pdf":
            if not content.startswith(b"%PDF-"):
                raise ExtractionError("The file does not have a valid PDF signature.")
            reader = PdfReader(io.BytesIO(content))
            if reader.is_encrypted:
                raise ExtractionError("Password-protected PDFs are not supported. Upload an unlocked copy.")
            pages = len(reader.pages)
            if pages > MAX_PAGES:
                raise ExtractionError("The maximum PDF length is 50 pages.")
            extracted = []
            for number, page in enumerate(reader.pages, 1):
                # A valid blank page can omit /Contents entirely.
                page_text = (page.extract_text(extraction_mode="layout") or "") if "/Contents" in page else ""
                if len(page.images):
                    raise ExtractionError(f"PDF page {number} needs OCR. Enable local page processing and confirm preview storage, then upload again; no partial document was accepted.")
                page_text = page_text.strip()
                start = len("\n\n".join(extracted)) + (2 if extracted and page_text else 0)
                details.append({"page": number, "method": "text_layer", "start": start, "end": start + len(page_text), "characters": len(page_text), "uncertain_regions": []})
                if page_text:
                    extracted.append(page_text)
            text = "\n\n".join(extracted)
            if not text.strip():
                raise ExtractionError("This PDF has no extractable text. Enable local OCR page processing and confirm preview storage, then upload again.")
        elif extension in {".docx", ".xlsx"}:
            if not content.startswith(b"PK"):
                raise ExtractionError("The file is not a valid Office document.")
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if len(archive.infolist()) > 2000 or sum(item.file_size for item in archive.infolist()) > 40 * 1024 * 1024:
                    raise ExtractionError("The Office document expands beyond the safe processing limit.")
                required = "word/document.xml" if extension == ".docx" else "xl/workbook.xml"
                if required not in archive.namelist():
                    raise ExtractionError("The file contents do not match its extension.")
                if any(name.startswith(("word/media/", "xl/media/")) for name in archive.namelist()):
                    raise ExtractionError("Embedded Office images are not processed in this stage; no partial document was accepted.")
            if extension == ".docx":
                document = Document(io.BytesIO(content))
                blocks = []
                for block in document.iter_inner_content():
                    if hasattr(block, "rows"):
                        blocks.extend(" | ".join(cell.text for cell in row.cells) for row in block.rows)
                    else:
                        blocks.append(block.text)
                text = "\n".join(blocks)
            else:
                workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
                blocks = []
                cells = 0
                text_bytes = 0
                try:
                    for sheet in workbook:
                        heading=f"Sheet: {sheet.title}"
                        text_bytes += len(heading.encode("utf-8")) + int(bool(blocks))
                        blocks.append(heading)
                        for row in sheet.iter_rows():
                            cells += len(row)
                            if cells > 100_000:
                                raise ExtractionError("The spreadsheet exceeds the 100,000-cell processing limit.")
                            if any(cell.data_type == "f" for cell in row):
                                raise ExtractionError("This stage requires spreadsheet values, not formulas. Export a values-only copy.")
                            values = [str(cell.value) if cell.value is not None else "" for cell in row]
                            if any(values):
                                line=" | ".join(values)
                                text_bytes += len(line.encode("utf-8")) + 1
                                blocks.append(line)
                            if text_bytes > MAX_TEXT:
                                raise ExtractionError("The spreadsheet contains too much text for this local extraction stage.")
                finally:
                    workbook.close()
                text = "\n".join(blocks)
        else:
            raise ExtractionError("Supported files at this stage: TXT, text-based PDF, DOCX and XLSX.")
    except ExtractionError:
        raise
    except Exception:
        raise ExtractionError("The document could not be read. It may be corrupted or use an unsupported format.") from None
    return {"name": Path(name).name, "size": len(content), "pages": pages, "text": validate_text(text), "page_details": details}
