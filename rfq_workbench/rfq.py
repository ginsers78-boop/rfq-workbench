"""Generate immutable local RFQ packages; no mailbox integration or sending."""
import hashlib
import io
import json
import re
import uuid
import zipfile
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import BaseModel, ConfigDict, Field

from .extraction import ExtractionError
from .privacy import check_bundle, check_text, suggested_markers

HEADERS = ["RFQ Reference", "Line ID", "Original Line", "Description", "Specification", "Quantity", "Unit",
           "Unit Price", "Currency", "Available Quantity", "MOQ", "Lead Time (days)", "Brand", "Country of Origin", "Validity Date", "Notes"]
PRODUCT_KEYS = ("id", "line_number", "description", "specification", "quantity", "unit")


class PrivacyInput(BaseModel):
    model_config = ConfigDict(extra="forbid", strict=True)
    revision: int = Field(ge=1)
    markers: list[str] = Field(max_length=100)
    confirmed: bool


class RFQInput(BaseModel):
    model_config = ConfigDict(extra="forbid", strict=True)
    revision: int = Field(ge=1)
    selection: dict[str, list[str]]


def save_privacy(store, identifier, payload):
    try:
        values = PrivacyInput.model_validate(payload)
    except ValueError:
        raise ExtractionError("The privacy checklist is invalid.") from None
    if not values.confirmed:
        raise ExtractionError("Review the source and confirm the confidential-data checklist first.")
    if any(not value.strip() or len(value) > 500 for value in values.markers):
        raise ExtractionError("Privacy markers must be nonempty and at most 500 characters each.")
    record = store.get(identifier)
    markers = list(dict.fromkeys(suggested_markers(record["document"]["text"]) + [value.strip() for value in values.markers]))
    record["privacy"] = {"markers": markers, "confirmed": True}
    record["revision"] = values.revision + 1
    with store.connect() as connection:
        changed = connection.execute("UPDATE requirements SET revision = ?, payload = ? WHERE id = ? AND revision = ?", (record["revision"], json.dumps(record), identifier, values.revision)).rowcount
        if changed != 1:
            raise ExtractionError("The requirement changed. Reopen it before saving the privacy checklist.")
        store.record_history(connection, record, "privacy_confirmed")
    return record


def reply_workbook(reference, lines):
    book = Workbook()
    book.properties.creator = "RFQ Workbench"
    sheet = book.active
    sheet.title = "Supplier Reply"
    sheet.append(HEADERS)
    for line in lines:
        sheet.append([reference, line["id"], line["line_number"], line["description"], line["specification"] or "", line["quantity"], line["unit"]] + [None] * 9)
    for row in sheet:
        for cell in row:
            if isinstance(cell.value, str):
                # All document-derived strings are literal text, including leading '='.
                cell.data_type = "s"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if cell.row == 1:
                cell.fill = PatternFill("solid", fgColor="215C4C")
                cell.font = Font(color="FFFFFF", bold=True)
            elif cell.column >= 8:
                cell.fill = PatternFill("solid", fgColor="E9F2FA")
                cell.font = Font(color="175C97")
        if row[0].row > 1:
            sheet.row_dimensions[row[0].row].height = 60
    widths = [32, 40, 15, 36, 44, 14, 14, 16, 14, 20, 14, 20, 22, 24, 19, 40]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = width
    sheet.row_dimensions[1].height = 34
    sheet.freeze_panes = "H2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.zoomScale = 80
    sheet.print_title_rows = "1:1"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    for column in ["H", "J", "K", "L"]:
        rule = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1=0, allow_blank=True)
        rule.errorTitle = "Enter a nonnegative number"
        rule.error = "Use a number of zero or greater, or leave the field blank."
        rule.showErrorMessage = True
        sheet.add_data_validation(rule)
        rule.add(f"{column}2:{column}{len(lines) + 1}")
    for row in range(2, len(lines) + 2):
        sheet.cell(row, 8).number_format = "0.0000"
        sheet.cell(row, 15).number_format = "yyyy-mm-dd"
    instructions = book.create_sheet("Instructions")
    messages = ["Supplier reply instructions", f"RFQ reference: {reference}",
                "Enter your offer in the blue columns on Supplier Reply. Leave unknown values blank.",
                "Keep the RFQ Reference, Line ID and Original Line unchanged so replies can be matched.",
                "Use one currency code per offered line, e.g. OMR, AED, USD or EUR. Enter unit price, not a line total.",
                "Enter lead time in days and validity as YYYY-MM-DD. Use Notes for alternatives or qualifications.",
                "Include your company name in your email reply and return the completed workbook.",
                "No customer enquiry or customer selling price is included in this template."]
    for message in messages:
        instructions.append([message])
    instructions.column_dimensions["A"].width = 110
    for row in instructions:
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
        instructions.row_dimensions[row[0].row].height = 36
    instructions["A1"].font = Font(size=16, bold=True, color="215C4C")
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def build_bundle(groups, markers):
    buffer = io.BytesIO()
    public = []
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for group in groups:
            reference = group["reference"]
            product_rows = []
            for line in group["lines"]:
                fields = [line["line_number"], line["description"], line["specification"] or "", f'{line["quantity"]:g}', line["unit"]]
                product_rows.append(" | ".join(re.sub(r"\s+", " ", str(field)).strip() for field in fields))
            subject = f"Request for quotation · {reference} · {group['category']}"
            body = (f"Dear supplier,\n\nPlease quote the following {group['category']} items under reference {reference}.\n"
                    "Please complete the attached Excel reply template and keep its reference and line IDs unchanged.\n\n"
                    "Original line | Description | Specification | Quantity | Unit\n" + "\n".join(product_rows) +
                    "\n\nPlease state unit price, currency, available quantity, MOQ, lead time, brand, origin and offer validity.\n"
                    f"Include {reference} in your reply subject. Please identify your company in your reply.\n\nThank you.\n")
            message = {"reference": reference, "category": group["category"], "subject": subject, "body": body,
                       "bcc": [{"supplier_id": supplier["id"], "name": supplier["name"], "email": supplier["email"]} for supplier in group["suppliers"]],
                       "attachment": f"{reference}/reply.xlsx", "lines": group["lines"]}
            check_text(json.dumps(message, ensure_ascii=False), markers)
            archive.writestr(f"{reference}/email-body.txt", f"Subject: {subject}\n\n{body}")
            archive.writestr(f"{reference}/reply.xlsx", reply_workbook(reference, group["lines"]))
            public.append(message)
        archive.writestr("dispatch-manifest.json", json.dumps({"messages": public}, ensure_ascii=False, indent=2))
        archive.writestr("READ-ME.txt", "LOCAL RFQ PREPARATION PACKAGE - NO EMAIL WAS SENT OR CREATED IN GMAIL.\n\n"
                          "This ZIP and dispatch-manifest.json are INTERNAL: they contain the BCC supplier list.\n"
                          "For each category, place only the reviewed email body in your mail client and attach its reply.xlsx.\n"
                          "Put selected suppliers in BCC, never To/CC. Do not send the ZIP, manifest or original customer file.\n"
                          "Review the exact message, recipients and attachment before any dispatch.\n")
    bundle = buffer.getvalue()
    check_bundle(bundle, markers)
    return public, bundle


def create_rfq(store, identifier, payload):
    try:
        values = RFQInput.model_validate(payload)
    except ValueError:
        raise ExtractionError("Select suppliers by category and use the current saved requirement revision.") from None
    with store.connect() as connection:
        connection.execute("BEGIN IMMEDIATE")
        row = connection.execute("SELECT payload FROM requirements WHERE id = ?", (identifier,)).fetchone()
        if not row:
            raise ExtractionError("The requirement was not found.")
        record = json.loads(row[0])
        if record["revision"] != values.revision:
            raise ExtractionError("The requirement changed. Reopen it before preparing RFQs.")
        if record["status"] != "reviewed":
            raise ExtractionError("Review all included lines or exclude incorrect lines with a reason before preparing RFQs.")
        if record["document"].get("used_ocr") and record.get("source_verified") is not True:
            raise ExtractionError("Compare the OCR transcript and product lines with the source page images first.")
        if len(record["analysis"].get("parts", [])) > 1 and record.get("coverage_verified") is not True:
            raise ExtractionError("Confirm review of the complete source and all model parts before preparing RFQs.")
        privacy = record.get("privacy", {})
        if privacy.get("confirmed") is not True:
            raise ExtractionError("Complete the confidential-data checklist first.")
        included = [line for line in record["lines"] if not line.get("excluded", False)]
        categories = sorted({line["category"] for line in included})
        if not categories or set(values.selection) != set(categories):
            raise ExtractionError("Select suppliers for every included category, without adding other categories.")
        groups, suppliers_snapshot = [], []
        for category in categories:
            ids = values.selection[category]
            if not 1 <= len(ids) <= 100 or len(set(ids)) != len(ids):
                raise ExtractionError("Choose one to 100 distinct suppliers per category.")
            suppliers = []
            for supplier_id in sorted(ids):
                row = connection.execute("SELECT payload FROM suppliers WHERE id = ?", (supplier_id,)).fetchone()
                if not row:
                    raise ExtractionError("A selected supplier was not found.")
                supplier = json.loads(row[0])
                if not supplier["active"] or category not in supplier["categories"]:
                    raise ExtractionError("A selected supplier is inactive or does not supply this category.")
                suppliers.append(supplier)
            suppliers_snapshot.extend(suppliers)
            lines = [{key: line[key] for key in PRODUCT_KEYS} for line in included if line["category"] == category]
            for line in lines:
                if not line["description"].strip() or not line["unit"] or line["quantity"] is None or line["quantity"] <= 0:
                    raise ExtractionError("Every RFQ line needs a product, positive quantity and unit.")
                if any(re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", str(value)) for value in line.values() if value is not None):
                    raise ExtractionError("Remove unsupported control characters from product fields.")
            groups.append({"category": category, "reference": "RFQ-" + datetime.now(timezone.utc).strftime("%Y%m%d") + "-" + uuid.uuid4().hex[:12].upper(), "suppliers": suppliers, "lines": lines})
        signature = hashlib.sha256(json.dumps({"requirement": identifier, "revision": values.revision, "suppliers": suppliers_snapshot, "selection": values.selection}, sort_keys=True).encode()).hexdigest()
        row = connection.execute("SELECT payload FROM rfqs WHERE signature = ?", (signature,)).fetchone()
        if row:
            return json.loads(row[0])
        markers = list(dict.fromkeys(suggested_markers(record["document"]["text"]) + privacy["markers"]))
        messages, bundle = build_bundle(groups, markers)
        result = {"id": str(uuid.uuid4()), "requirement_id": identifier, "revision": values.revision, "created": datetime.now(timezone.utc).isoformat(),
                  "status": "preview", "messages": messages, "suppliers_snapshot": suppliers_snapshot}
        connection.execute("INSERT INTO rfqs VALUES (?, ?, ?, ?, ?)", (result["id"], identifier, signature, json.dumps(result), bundle))
    return result


def current_rfq(connection, identifier):
    row = connection.execute("SELECT payload, bundle FROM rfqs WHERE id = ?", (identifier,)).fetchone()
    if not row:
        raise ExtractionError("The RFQ package was not found.")
    record = json.loads(row[0])
    parent = connection.execute("SELECT revision FROM requirements WHERE id = ?", (record["requirement_id"],)).fetchone()
    if not parent or parent[0] != record["revision"]:
        raise ExtractionError("This RFQ is outdated. Prepare a new package from the current requirement.")
    for supplier in record["suppliers_snapshot"]:
        revision = connection.execute("SELECT revision FROM suppliers WHERE id = ?", (supplier["id"],)).fetchone()
        if not revision or revision[0] != supplier["revision"]:
            raise ExtractionError("A supplier changed. Prepare a new RFQ package before downloading.")
    return record, row[1]


def approve_rfq(store, identifier):
    with store.connect() as connection:
        connection.execute("BEGIN IMMEDIATE")
        record, _ = current_rfq(connection, identifier)
        record["status"] = "approved_for_local_export"
        record["approved_at"] = datetime.now(timezone.utc).isoformat()
        connection.execute("UPDATE rfqs SET payload = ? WHERE id = ?", (json.dumps(record), identifier))
    return record


def download_rfq(store, identifier):
    with store.connect() as connection:
        connection.execute("BEGIN")
        record, bundle = current_rfq(connection, identifier)
        if record["status"] != "approved_for_local_export":
            raise ExtractionError("Inspect and approve the RFQ preview before downloading.")
        return bundle
